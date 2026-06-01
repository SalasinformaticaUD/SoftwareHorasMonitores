from __future__ import annotations

from dataclasses import dataclass
from datetime import datetime
from html import escape
from pathlib import Path
import logging
import re
import os
from io import BytesIO
from xml.sax.saxutils import escape
from reportlab.lib.utils import ImageReader
from pathlib import Path
from django.utils import timezone
 
from reportlab.lib import colors
from reportlab.lib.enums import TA_CENTER, TA_JUSTIFY, TA_LEFT
from reportlab.lib.pagesizes import letter
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import cm, mm
from reportlab.platypus import (
    HRFlowable,
    Image,
    Paragraph,
    SimpleDocTemplate,
    Spacer,
    Table,
    TableStyle,
)

from django.conf import settings
from django.core.exceptions import ValidationError
from django.core.files.base import ContentFile
from django.core.mail import EmailMessage
from django.db import transaction
from django.utils import timezone
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from reportlab.lib import colors
from reportlab.lib.enums import TA_CENTER, TA_JUSTIFY
from reportlab.lib.pagesizes import letter
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

from apps.common.choices import DepartmentChoices, SessionStateChoices
from apps.common.events import DomainEvent, event_bus
from apps.common.utils import normalize_text
from apps.reports.events import REPORT_GENERATED
from apps.reports.models import MonitorMemorandum, MonitorReportSnapshot
from apps.reports.selectors import aggregate_monitor_metrics, build_monitor_rows_for_user


logger = logging.getLogger(__name__)


@transaction.atomic
def generate_monitor_report(*, monitor, start_date, end_date, generated_by=None) -> MonitorReportSnapshot:
    metrics = aggregate_monitor_metrics(monitor=monitor, start_date=start_date, end_date=end_date)
    snapshot, _ = MonitorReportSnapshot.objects.update_or_create(
        monitor=monitor,
        start_date=start_date,
        end_date=end_date,
        defaults={
            "generated_by": generated_by,
            "department": monitor.department,
            "normal_minutes": metrics["normal_minutes"],
            "approved_overtime_minutes": metrics["approved_overtime_minutes"],
            "pending_overtime_minutes": metrics["pending_overtime_minutes"],
            "penalty_minutes": metrics["penalty_minutes"],
            "late_count": metrics["late_count"],
            "annotation_delta_minutes": metrics["annotation_delta_minutes"],
            "total_minutes": metrics["total_minutes"],
            "has_memorandum": metrics["has_memorandum"],
        },
    )
    event_bus.publish(
        DomainEvent(
            name=REPORT_GENERATED,
            aggregate_id=str(snapshot.id),
            payload={
                "report_id": str(snapshot.id),
                "department": snapshot.department,
                "monitor_id": str(snapshot.monitor_id),
            },
        )
    )
    return snapshot


DASHBOARD_EXPORT_DIRECTORY = Path(settings.MEDIA_ROOT) / "dashboard_exports"
DEPARTMENT_EXPORT_FILENAMES = {
    DepartmentChoices.INFORMATICS_LABS: "dashboard_monitores.xlsx",
    DepartmentChoices.PHYSICS: "dashboard_monitores_fisica.xlsx",
    DepartmentChoices.ELECTRICAL: "dashboard_monitores_laboratorios.xlsx",
}
SIGNED_COMMITMENT_ACTS_FOLDER = "actas_compromiso_firmadas"
MEMORANDUM_LATE_THRESHOLD = 3


@dataclass(frozen=True)
class CommitmentActStatus:
    """Estado administrativo de firma del acta de un monitor.

    Attributes:
        monitor: Monitor relacionado con la fila.
        signed_file: Ruta absoluta del PDF firmado cuando existe.
        signed_file_name: Nombre visible del archivo firmado.
        uploaded_at: Fecha detectada desde la ultima modificacion del archivo.
    """

    monitor: object
    signed_file: Path | None
    signed_file_name: str
    uploaded_at: datetime | None

    @property
    def has_signed(self) -> bool:
        """Indica si el monitor ya tiene un acta firmada en carpeta.

        Returns:
            bool: ``True`` si se encontro un PDF firmado para el monitor.
        """

        return self.signed_file is not None


def get_dashboard_export_directory() -> Path:
    DASHBOARD_EXPORT_DIRECTORY.mkdir(parents=True, exist_ok=True)
    return DASHBOARD_EXPORT_DIRECTORY


def get_signed_commitment_acts_directory() -> Path:
    """Retorna la carpeta donde se almacenan las actas firmadas.

    Returns:
        Path: Directorio de actas firmadas, creado si no existia.
    """

    directory = Path(settings.MEDIA_ROOT) / SIGNED_COMMITMENT_ACTS_FOLDER
    directory.mkdir(parents=True, exist_ok=True)
    return directory


def _monitor_code_pattern(monitor) -> re.Pattern[str]:
    code = re.escape(normalize_text(str(monitor.codigo_estudiante or "")))
    return re.compile(rf"(^|[_\-\s]){code}($|[_\-\s])")


def _signed_act_candidates() -> list[Path]:
    directory = get_signed_commitment_acts_directory()
    return [path for path in directory.iterdir() if path.is_file() and path.suffix.lower() == ".pdf"]


def signed_commitment_act_for_monitor(monitor) -> Path | None:
    """Busca el PDF firmado mas reciente asociado a un monitor.

    Args:
        monitor: Instancia de ``Monitor`` usada para cruzar por codigo estudiantil.

    Returns:
        Path | None: Ruta del PDF firmado mas reciente, o ``None`` si no existe.
    """

    if not monitor.codigo_estudiante:
        return None
    pattern = _monitor_code_pattern(monitor)
    matches = [
        path
        for path in _signed_act_candidates()
        if pattern.search(normalize_text(path.stem))
    ]
    if not matches:
        return None
    return max(matches, key=lambda path: path.stat().st_mtime)


def commitment_act_status_for_monitor(monitor) -> CommitmentActStatus:
    """Construye el estado de acta firmada para un monitor.

    Args:
        monitor: Instancia de ``Monitor`` a evaluar.

    Returns:
        CommitmentActStatus: Datos de estado y archivo firmado detectado.
    """

    signed_file = signed_commitment_act_for_monitor(monitor)
    uploaded_at = None
    if signed_file:
        uploaded_at = timezone.localtime(datetime.fromtimestamp(signed_file.stat().st_mtime, tz=timezone.get_current_timezone()))
    return CommitmentActStatus(
        monitor=monitor,
        signed_file=signed_file,
        signed_file_name=signed_file.name if signed_file else "",
        uploaded_at=uploaded_at,
    )


def build_commitment_act_status_rows(monitors) -> list[CommitmentActStatus]:
    """Cruza una coleccion de monitores con los PDFs firmados disponibles.

    Args:
        monitors: Iterable o QuerySet de monitores visibles para el usuario.

    Returns:
        list[CommitmentActStatus]: Filas listas para mostrar en el modulo admin.
    """

    return [commitment_act_status_for_monitor(monitor) for monitor in monitors]

 
LOGO_PATH = r"C:\Users\ud\Documents\MonitoresV1.1.0\SoftwareHorasMonitores\static\branding\logo-ud.png"
MEMORANDUM_LATE_THRESHOLD = 3
 
 
def _monitor_email(monitor) -> str:
    user = getattr(monitor, "user", None)
    return (getattr(user, "email", "") or "").strip()


def _memorandum_filename(*, monitor, late_count: int) -> str:
    safe_name = normalize_text(monitor.full_name).replace(" ", "_") or "monitor"
    return f"Memorando_{late_count}_retardos_{safe_name}_{monitor.codigo_estudiante}.pdf"


def _format_time(value) -> str:
    value = _as_time(value)
    return value.strftime("%H:%M:%S") if getattr(value, "second", 0) else value.strftime("%H:%M")


def _as_time(value):
    if isinstance(value, datetime):
        if timezone.is_aware(value):
            value = timezone.localtime(value)
        return value.time()
    return value


def _lateness_observation(session) -> str:
    schedule = session.schedule
    actual_start = _as_time(session.actual_start)
    if not schedule:
        return f"LLEGO A LAS {_format_time(actual_start)}"

    scheduled_at = datetime.combine(session.work_day, schedule.start_time)
    actual_at = datetime.combine(session.work_day, actual_start)
    late_seconds = max(int((actual_at - scheduled_at).total_seconds()), 0)
    minutes, seconds = divmod(late_seconds, 60)
    if seconds:
        delay = f"{minutes} MIN {seconds} SEG"
    else:
        delay = f"{minutes} MINUTOS"
    return f"LLEGO A LAS {_format_time(actual_start)} ({delay} TARDE)"


def _lateness_sessions_for_memorandum(*, monitor, late_count: int):
    """Retorna solo retardos reales para el detalle del memorando."""
    from apps.work_sessions.models import WorkSession

    sessions = []
    candidates = (
        WorkSession.objects.select_related("monitor", "schedule")
        .filter(monitor=monitor, is_late=True)
        .exclude(session_state=SessionStateChoices.INVALID)
        .order_by("-work_day", "-actual_start")
    )
    for session in candidates.iterator():
        if not _counts_for_lateness_memorandum(session):
            continue
        sessions.append(session)
        if len(sessions) >= late_count:
            break
    sessions.reverse()
    return sessions


def _unexcused_lateness_count(*, monitor) -> int:
    from apps.work_sessions.models import WorkSession

    candidates = (
        WorkSession.objects.select_related("monitor")
        .filter(monitor=monitor, is_late=True, lateness_excused=False)
        .exclude(session_state=SessionStateChoices.INVALID)
    )
    return sum(1 for session in candidates.iterator() if _counts_for_lateness_memorandum(session))


def _counts_for_lateness_memorandum(session) -> bool:
    from apps.schedules.selectors import lateness_exception_for

    return (
        not session.lateness_excused
        and lateness_exception_for(monitor=session.monitor, day=session.work_day) is None
    )


def _deliver_lateness_memorandum_email(
    *,
    monitor,
    email: str,
    filename: str,
    pdf_bytes: bytes,
    late_count: int,
    updated: bool = False,
) -> None:
    qualifier = "actualizado " if updated else ""
    message = EmailMessage(
        subject=f"Memorando por {late_count} llegadas tarde",
        body=(
            f"Cordial saludo {monitor.full_name},\n\n"
            f"Adjuntamos el memorando {qualifier}generado por completar "
            f"{late_count} llegadas tarde acumuladas.\n\n"
            "Sistema de Registro de Asistencia"
        ),
        from_email=getattr(settings, "DEFAULT_FROM_EMAIL", None),
        to=[email],
    )
    message.attach(filename, pdf_bytes, "application/pdf")
    message.send(fail_silently=False)


def refresh_lateness_memorandum_pdf(*, memorandum: MonitorMemorandum) -> tuple[str, bytes]:
    """Regenera el PDF de un memorando usando las excepciones vigentes."""

    monitor = memorandum.monitor
    late_count = memorandum.late_count_threshold
    filename = _memorandum_filename(monitor=monitor, late_count=late_count)
    pdf_bytes = generate_lateness_memorandum_pdf(monitor=monitor, late_count=late_count)
    memorandum.pdf_file.save(filename, ContentFile(pdf_bytes), save=False)
    memorandum.save(update_fields=["pdf_file", "updated_at"])
    return filename, pdf_bytes


def _load_logo() -> ImageReader | None:
    """Carga el logo como ImageReader en memoria usando Pillow.
    Compone sobre fondo blanco para evitar el fondo negro en ReportLab."""
    if not os.path.exists(LOGO_PATH):
        return None
    try:
        from PIL import Image as PILImage
        img = PILImage.open(LOGO_PATH)
        # Convertir a RGBA primero si es modo P (paleta)
        if img.mode == "P":
            img = img.convert("RGBA")
        # Componer sobre fondo blanco para eliminar transparencia
        if img.mode in ("RGBA", "LA"):
            background = PILImage.new("RGB", img.size, (255, 255, 255))
            background.paste(img, mask=img.split()[-1])
            img = background
        else:
            img = img.convert("RGB")
        buf = BytesIO()
        img.save(buf, format="PNG")
        buf.seek(0)
        return ImageReader(buf)
    except Exception:
        return None
 
 
def generate_lateness_memorandum_pdf(*, monitor, late_count: int) -> bytes:
    """Genera el PDF de memorando por llegadas tarde con el diseño oficial
    de la Universidad Distrital Francisco José de Caldas."""
    buffer = BytesIO()
 
    LEFT   = 1.8 * cm
    RIGHT  = 1.8 * cm
    TOP    = 3.8 * cm
    BOTTOM = 2.5 * cm
 
    document = SimpleDocTemplate(
        buffer,
        pagesize=letter,
        title="Memorando",
        leftMargin=LEFT,
        rightMargin=RIGHT,
        topMargin=TOP,
        bottomMargin=BOTTOM,
    )
 
    page_width = letter[0] - LEFT - RIGHT
 
    # ── Estilos ──────────────────────────────────────────────────────────────
    styles = getSampleStyleSheet()
 
    normal = ParagraphStyle(
        "MemoNormal",
        parent=styles["Normal"],
        fontName="Helvetica",
        fontSize=10,
        leading=14,
        spaceAfter=2,
    )
    bold_label = ParagraphStyle(
        "MemoBoldLabel",
        parent=normal,
        fontName="Helvetica-Bold",
    )
    justified = ParagraphStyle(
        "MemoJustified",
        parent=normal,
        alignment=TA_JUSTIFY,
    )
    centered_title = ParagraphStyle(
        "MemoTitle",
        parent=styles["Normal"],
        fontName="Helvetica-Bold",
        fontSize=11,
        leading=14,
        alignment=TA_CENTER,
        spaceAfter=6,
    )
    th = ParagraphStyle(
        "MemoTH",
        parent=styles["Normal"],
        fontName="Helvetica-Bold",
        fontSize=8,
        leading=10,
        alignment=TA_CENTER,
    )
    td = ParagraphStyle(
        "MemoTD",
        parent=styles["Normal"],
        fontName="Helvetica",
        fontSize=8,
        leading=10,
        alignment=TA_LEFT,
    )
    td_center = ParagraphStyle(
        "MemoTDCenter",
        parent=td,
        alignment=TA_CENTER,
    )
 
    # ── Datos ────────────────────────────────────────────────────────────────
    memorandum_number = late_count // MEMORANDUM_LATE_THRESHOLD
    current_date = timezone.localdate()
 
    sessions = _lateness_sessions_for_memorandum(monitor=monitor, late_count=late_count)
 
    def _session_subject(session) -> str:
        schedule = session.schedule
        if schedule and schedule.asignatura:
            return schedule.asignatura
        if schedule and schedule.location:
            return schedule.location
        return monitor.get_department_display()
 
    def _session_hour(session) -> str:
        schedule = session.schedule
        start = schedule.start_time if schedule else session.actual_start
        end   = schedule.end_time   if schedule else session.actual_end
        return f"{start.strftime('%H:%M')} - {end.strftime('%H:%M')}"
 
    # ── Tabla de retardos ────────────────────────────────────────────────────
    col_widths = [
        page_width * 0.38,
        page_width * 0.15,
        page_width * 0.16,
        page_width * 0.31,
    ]
 
    rows = [[
        Paragraph("<b>NOMBRE DE LA CLASE O LABORATORIO</b>", th),
        Paragraph("<b>HORA</b>", th),
        Paragraph("<b>FECHA D/M/A</b>", th),
        Paragraph("<b>OBSERVACIONES</b>", th),
    ]]
 
    if sessions:
        for session in sessions:
            rows.append([
                Paragraph(escape(_session_subject(session)), td),
                Paragraph(escape(_session_hour(session)), td_center),
                Paragraph(session.work_day.strftime("%d/%m/%Y"), td_center),
                Paragraph(escape(_lateness_observation(session)), td),
            ])
    else:
        rows.append([
            Paragraph("Monitoria asignada", td),
            Paragraph("-", td_center),
            Paragraph("-", td_center),
            Paragraph(f"{late_count} llegadas tarde acumuladas", td),
        ])
 
    details_table = Table(rows, colWidths=col_widths, hAlign="LEFT")
    details_table.setStyle(TableStyle([
        ("BACKGROUND",    (0, 0), (-1, 0),  colors.HexColor("#E8E8E8")),
        ("FONTNAME",      (0, 0), (-1, 0),  "Helvetica-Bold"),
        ("GRID",          (0, 0), (-1, -1), 0.5, colors.black),
        ("LEFTPADDING",   (0, 0), (-1, -1), 5),
        ("RIGHTPADDING",  (0, 0), (-1, -1), 5),
        ("TOPPADDING",    (0, 0), (-1, -1), 4),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 4),
        ("VALIGN",        (0, 0), (-1, -1), "MIDDLE"),
    ]))
 
    # ── Encabezado y pie (canvas) ────────────────────────────────────────────
    # El logo se carga UNA sola vez fuera del callback para no repetir I/O
    logo_reader = _load_logo()
 
    def _draw_header_footer(canvas, doc):
        canvas.saveState()
        page_w, page_h = letter
 
        # --- Logo ---
        logo_x = LEFT
        logo_y = page_h - 2.7 * cm
        logo_w = 3.0 * cm
        logo_h = 2.2 * cm
 
        if logo_reader is not None:
            canvas.drawImage(
                logo_reader,
                logo_x, logo_y,
                width=200, height=logo_h,
                preserveAspectRatio=True,
                anchor="nw",
            )
 
       
 
        # --- Línea separadora ---
        canvas.setStrokeColor(colors.black)
        canvas.setLineWidth(0.5)
        canvas.line(LEFT, page_h - 2.85 * cm, page_w - RIGHT, page_h - 2.85 * cm)
 
        # --- Acreditación ---
        canvas.setFont("Helvetica", 6)
        canvas.drawString(
            LEFT, page_h - 3.2 * cm,
            "Acreditación Institucional de Alta Calidad. Resolución No. 23096 del 15 de diciembre de 2016",
        )
        canvas.drawRightString(page_w - RIGHT, page_h - 3.2 * cm, "labiud@udistrital.edu.co")
 
        # --- Pie de página ---
        footer_y = 1.6 * cm
        canvas.setLineWidth(0.5)
        canvas.line(LEFT, footer_y + 0.7 * cm, page_w - RIGHT, footer_y + 0.7 * cm)
 
        canvas.setFont("Helvetica", 7)
        canvas.drawString(LEFT, footer_y + 0.35 * cm, "PBX 57(1)323 9300 Exts. 1520 – 1521 - 1525")
        canvas.drawString(
            LEFT, footer_y,
            "Carrera 8 No 40 62, Piso 5, Edificio Sabio Caldas, Bogotá D.C. – Colombia",
        )
        canvas.drawRightString(page_w - RIGHT, footer_y + 0.35 * cm, "Línea de atención gratuita")
        canvas.drawRightString(page_w - RIGHT, footer_y, "01 800 091 44 10")
 
        canvas.setFont("Helvetica", 8)
        canvas.drawCentredString(page_w / 2, footer_y - 0.45 * cm, str(doc.page))
 
        canvas.restoreState()
 
    # ── Story ────────────────────────────────────────────────────────────────
    story = []
 
    story.append(Paragraph(
        f"<b>MEMORANDO N. {memorandum_number:03d}-{current_date.year}</b>",
        centered_title,
    ))
    story.append(Spacer(1, 8))
 
    label_w = 2.0 * cm
    value_w = page_width - label_w
 
    def _meta_row(label: str, value: str):
        return [
            Paragraph(f"<b>{label}</b>", bold_label),
            Paragraph(value, normal),
        ]
 
    meta_rows = [
        _meta_row("DE:", "COORDINADOR DE LABORATORIOS-FACULTAD DE INGENIERÍA"),
        [Paragraph("", normal), Paragraph("Ing. EDILBERTO SUÁREZ TORRES", bold_label)],
        _meta_row("PARA:", f"{escape(monitor.full_name)} - <b>Código: {escape(monitor.codigo_estudiante)}</b>"),
        [Paragraph("", normal), Paragraph(f"Monitor – {escape(monitor.get_department_display())}", normal)],
        _meta_row("ASUNTO:", "LLAMADO DE ATENCIÓN"),
        _meta_row("FECHA:", current_date.strftime("%d/%m/%Y")),
    ]
 
    meta_table = Table(meta_rows, colWidths=[label_w, value_w], hAlign="LEFT")
    meta_table.setStyle(TableStyle([
        ("VALIGN",        (0, 0), (-1, -1), "TOP"),
        ("LEFTPADDING",   (0, 0), (-1, -1), 0),
        ("RIGHTPADDING",  (0, 0), (-1, -1), 0),
        ("TOPPADDING",    (0, 0), (-1, -1), 1),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 1),
    ]))
    story.append(meta_table)
 
    story.append(Spacer(1, 4))
    story.append(HRFlowable(width="100%", thickness=0.5, color=colors.black))
    story.append(Spacer(1, 10))
 
    if memorandum_number == 1:
        story.append(Paragraph(
        "Por medio del presente se le hace el primer llamado de atención debido al "
        "incumplimiento en algunas tareas asignadas, recuerde que la puntualidad y "
        "cumplimiento de cada una de estas hacen parte de su compromiso como monitor "
        "que es brindar una atención eficaz a docentes y estudiantes del Laboratorio "
        "de la Facultad de Ingeniería.",
        justified,)
        )
    elif memorandum_number == 2:
        story.append(Paragraph(
        "Por medio del presente se le hace el segundo llamado de atención debido al "
        "incumplimiento en algunas tareas asignadas, recuerde que la puntualidad y "
        "cumplimiento de cada una de estas hacen parte de su compromiso como monitor "
        "que es brindar una atención eficaz a docentes y estudiantes del Laboratorio "
        "de la Facultad de Ingeniería.",
        justified,)
        )
    else:
        story.append(Paragraph(
        "Por medio del presente se le hace el tercer llamado de atención debido al "
        "incumplimiento en algunas tareas asignadas, en este momento debe acercarse al almac{en para resolver la situacion, "
        "recuerde que la puntualidad y "
        "cumplimiento de cada una de estas hacen parte de su compromiso como monitor "
        "que es brindar una atención eficaz a docentes y estudiantes del Laboratorio "
        "de la Facultad de Ingeniería.",
        justified,)
        )


    story.append(Spacer(1, 8))
    story.append(Paragraph(
        "A continuación, relaciono las fechas de sus retardos, inasistencias y/o "
        "faltas en las Monitorias asignadas:",
        justified,
    ))
    story.append(Spacer(1, 10))
 
    story.append(details_table)
    story.append(Spacer(1, 10))
 
    story.append(Paragraph(
        "Es importante que tenga en cuenta que, si persiste con el incumplimiento, "
        "se continuara con el debido proceso.",
        justified,
    ))
    story.append(Spacer(1, 16))
    story.append(Paragraph("Cordialmente,", normal))
    story.append(Spacer(1, 36))
 
    story.append(Paragraph("<b>ING. EDILBERTO SUÁREZ TORRES</b>", normal))
    story.append(Paragraph("Coordinador Laboratorios", normal))
    story.append(Paragraph("Facultad de Ingeniería", normal))
    story.append(Spacer(1, 14))
 
    story.append(Paragraph("<b>NOTA:</b>", normal))
    story.append(Paragraph(
        "✓  LAS LLEGADAS TARDES, INASISTENCIAS Y FALLAS SE ACUMULAN Y AL FINAL DE "
        "CLASES DEBERA REALIZAR LA REPOSICIÓN LAS CUALES SERÁN ASIGNADAS POR EL "
        "PERSONAL DEL ALMACEN DE LABORATORIO – FAC. DE INGENIERIA",
        ParagraphStyle("MemoNota", parent=normal, fontSize=8, leading=11),
    ))
 
    document.build(story, onFirstPage=_draw_header_footer, onLaterPages=_draw_header_footer)
    return buffer.getvalue()


@transaction.atomic
def send_lateness_memorandum(*, memorandum: MonitorMemorandum) -> MonitorMemorandum:
    """Envia o reenvia un memorando existente al correo actual del monitor."""

    monitor = memorandum.monitor
    email = _monitor_email(monitor)
    if not email:
        raise ValidationError("El monitor no tiene correo registrado para enviar el memorando.")

    late_count = memorandum.late_count_threshold
    filename, pdf_bytes = refresh_lateness_memorandum_pdf(memorandum=memorandum)

    _deliver_lateness_memorandum_email(
        monitor=monitor,
        email=email,
        filename=filename,
        pdf_bytes=pdf_bytes,
        late_count=late_count,
        updated=memorandum.sent_at is not None,
    )
    memorandum.sent_to = email
    memorandum.sent_at = timezone.now()
    memorandum.save(update_fields=["sent_to", "sent_at", "updated_at"])
    return memorandum


@transaction.atomic
def create_and_send_lateness_memorandum(*, monitor, late_count: int) -> MonitorMemorandum | None:
    """Crea un memorando cuando se completa un bloque de tres retardos.

    Si el monitor tiene correo, tambien se envia. Si no tiene correo, se guarda
    el PDF y queda pendiente para reenviarlo cuando se complete el dato.

    Args:
        monitor: Monitor evaluado.
        late_count: Total actual de llegadas tarde.

    Returns:
        MonitorMemorandum | None: Memorando creado, o ``None`` si no aplica.
    """

    late_count = _unexcused_lateness_count(monitor=monitor)
    if late_count < MEMORANDUM_LATE_THRESHOLD or late_count % MEMORANDUM_LATE_THRESHOLD != 0:
        return None

    email = _monitor_email(monitor)
    
    memorandum, created = MonitorMemorandum.objects.get_or_create(
        monitor=monitor,
        late_count_threshold=late_count,
        defaults={"sent_to": email},
    )
    if not created:
        refresh_lateness_memorandum_pdf(memorandum=memorandum)
        return None

    pdf_bytes = generate_lateness_memorandum_pdf(monitor=monitor, late_count=late_count)
    filename = _memorandum_filename(monitor=monitor, late_count=late_count)
    memorandum.pdf_file.save(filename, ContentFile(pdf_bytes), save=False)
    memorandum.save(update_fields=["sent_to", "pdf_file", "updated_at"])

    if not email:
        return memorandum

    try:
        _deliver_lateness_memorandum_email(
            monitor=monitor,
            email=email,
            filename=filename,
            pdf_bytes=pdf_bytes,
            late_count=late_count,
        )
    except Exception:
        logger.exception(
            "No se pudo enviar el memorando de retardos para monitor %s.",
            monitor.id,
        )
        return memorandum
    memorandum.sent_to = email
    memorandum.sent_at = timezone.now()
    memorandum.save(update_fields=["sent_to", "sent_at", "updated_at"])
    return memorandum


def export_department_dashboard_to_excel(*, user, department: str) -> Path:
    export_directory = get_dashboard_export_directory()
    file_name = DEPARTMENT_EXPORT_FILENAMES.get(department, f"dashboard_{department}.xlsx")
    export_path = export_directory / file_name
    rows = build_monitor_rows_for_user(user, department=department)

    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "Dashboard"
    headers = [
        "Monitor",
        "Normales (h)",
        "Horas extra aprobadas (h)",
        "Horas extra por aprobar (h)",
        "Anotaciones (h)",
        "Total (h)",
        "Faltan para 192 h",
    ]
    worksheet.append(headers)

    header_fill = PatternFill(fill_type="solid", fgColor="D9EAF7")
    header_font = Font(bold=True)
    for column, header in enumerate(headers, start=1):
        cell = worksheet.cell(row=1, column=column, value=header)
        cell.font = header_font
        cell.fill = header_fill

    for row in rows:
        worksheet.append(
            [
                row["monitor"].full_name,
                row["normal_hours"],
                row["approved_overtime_hours"],
                row["pending_overtime_hours"],
                row["annotation_hours"],
                row["total_hours"],
                row["remaining_hours"],
            ]
        )

    widths = {
        "A": 36,
        "B": 16,
        "C": 24,
        "D": 24,
        "E": 18,
        "F": 14,
        "G": 20,
    }
    for column_letter, width in widths.items():
        worksheet.column_dimensions[column_letter].width = width
    worksheet.freeze_panes = "A2"

    workbook.save(export_path)
    workbook.close()
    return export_path
