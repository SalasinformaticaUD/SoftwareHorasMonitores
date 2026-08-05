import re
from dataclasses import dataclass, field
from datetime import datetime, time
from typing import Any, Iterable, Optional

from django.core.exceptions import ValidationError
from django.db import transaction
from openpyxl import load_workbook

from apps.attendance.validators import validate_excel_extension
from apps.common.choices import UserRoleChoices
from apps.common.utils import normalize_text
from apps.monitors.models import Monitor, PROJECT_CHOICES
from apps.monitors.selectors import visible_monitors_for_user
from apps.schedules.models import Schedule, ScheduleException
from apps.work_sessions.services import sync_sessions_for_exception_change


DAY_NAME_TO_WEEKDAY = {
    "lunes": Schedule.Weekday.MONDAY,
    "martes": Schedule.Weekday.TUESDAY,
    "miercoles": Schedule.Weekday.WEDNESDAY,
    "mircoles": Schedule.Weekday.WEDNESDAY,
    "jueves": Schedule.Weekday.THURSDAY,
    "viernes": Schedule.Weekday.FRIDAY,
    "sabado": Schedule.Weekday.SATURDAY,
    "sbado": Schedule.Weekday.SATURDAY,
    "domingo": Schedule.Weekday.SUNDAY,
}

DAY_HOUR_PATTERN = re.compile(
    r"(?P<day>[A-ZÁÉÍÓÚÜÑ]+)\s+(?P<start_hour>\d{1,2})(?::(?P<start_minute>\d{2}))?-(?P<end_hour>\d{1,2})(?::(?P<end_minute>\d{2}))?",
    re.IGNORECASE,
)


SCHEDULE_ROW_REQUIRED_COLUMNS = ("email_monitor", "day", "start_time", "end_time", "location")
SCHEDULE_ROW_COLUMN_ALIASES = {
    "email_monitor": (
        "email_monitor",
        "monitor email",
        "email monitor",
        "correo monitor",
        "correo del monitor",
        "email del monitor",
    ),
    "day": (
        "day",
        "weekday",
        "dia",
        "dia semana",
        "dia de la semana",
    ),
    "start_time": (
        "start_time",
        "start time",
        "start",
        "hora inicio",
        "hora inicial",
        "inicio",
    ),
    "end_time": (
        "end_time",
        "end time",
        "end",
        "hora fin",
        "hora final",
        "fin",
    ),
    "location": (
        "location",
        "place",
        "ubicacion",
        "lugar",
        "salon",
        "laboratorio",
    ),
    "asignatura": (
        "asignatura",
        "subject",
        "course",
        "materia",
    ),
    "grupo": (
        "grupo",
        "group",
    ),
    "docente": (
        "docente",
        "teacher",
        "profesor",
        "instructor",
    ),
    "proyecto_curricular": (
        "proyecto_curricular",
        "proyecto curricular",
        "programa",
        "curricular project",
        "academic program",
    ),
}
BLOCK_MONITOR_NAME_LABELS = ("NOMBRE", "NAME")
BLOCK_MONITOR_CODE_LABELS = ("CODIGO", "CODE", "STUDENT CODE")
BLOCK_DAY_HOUR_HEADER_ALIASES = {"dia hora", "day hour", "day time"}


@dataclass
class ParsedScheduleBlock:
    weekday: int
    start_time: time
    end_time: time


@dataclass
class ScheduleImportResult:
    created: int = 0
    reactivated: int = 0
    processed_monitors: int = 0
    skipped_rows: int = 0
    missing_monitors: list[str] = field(default_factory=list)
    unauthorized_monitors: list[str] = field(default_factory=list)


@dataclass
class ScheduleRowIssue:
    row_number: int
    monitor_email: str
    reason: str


@dataclass
class ScheduleBulkImportResult:
    total_rows: int = 0
    created: int = 0
    skipped: list[ScheduleRowIssue] = field(default_factory=list)
    errors: list[ScheduleRowIssue] = field(default_factory=list)


def _row_label_value(row_values: list[str], target_label: str) -> Optional[str]:
    normalized_target = normalize_text(target_label).rstrip(":")
    for index, value in enumerate(row_values):
        normalized_value = normalize_text(value).rstrip(":")
        if normalized_value == normalized_target:
            for candidate in row_values[index + 1 :]:
                if candidate:
                    return candidate.strip()
    return None


def _row_label_value_any(row_values: list[str], target_labels: Iterable[str]) -> Optional[str]:
    for label in target_labels:
        value = _row_label_value(row_values, label)
        if value is not None:
            return value
    return None


def _build_time(hour_text: str, minute_text: Optional[str]) -> time:
    return time(hour=int(hour_text), minute=int(minute_text or 0))


def _parse_day_hour(value: str) -> Optional[ParsedScheduleBlock]:
    normalized_value = (value or "").strip().replace("?", "").replace("�", "")
    match = DAY_HOUR_PATTERN.search(normalized_value)
    if not match:
        return None
    normalized_day = normalize_text(match.group("day")).replace("?", "").replace("�", "")
    weekday = DAY_NAME_TO_WEEKDAY.get(normalized_day)
    if weekday is None:
        return None
    start_time = _build_time(match.group("start_hour"), match.group("start_minute"))
    end_time = _build_time(match.group("end_hour"), match.group("end_minute"))
    if end_time <= start_time:
        raise ValidationError("La hora final del bloque debe ser posterior a la hora inicial.")
    return ParsedScheduleBlock(weekday=weekday, start_time=start_time, end_time=end_time)


def _merge_schedule_blocks(blocks: Iterable[ParsedScheduleBlock]) -> list[ParsedScheduleBlock]:
    ordered_blocks = sorted(blocks, key=lambda block: (block.weekday, block.start_time, block.end_time))
    merged: list[ParsedScheduleBlock] = []
    for block in ordered_blocks:
        if not merged:
            merged.append(block)
            continue
        previous = merged[-1]
        if previous.weekday == block.weekday and block.start_time <= previous.end_time:
            merged[-1] = ParsedScheduleBlock(
                weekday=previous.weekday,
                start_time=previous.start_time,
                end_time=max(previous.end_time, block.end_time),
            )
            continue
        merged.append(block)
    return merged


def _save_schedule_block(*, monitor: Monitor, block: ParsedScheduleBlock, result: ScheduleImportResult) -> None:
    schedule = Schedule.objects.filter(
        monitor=monitor,
        weekday=block.weekday,
        start_time=block.start_time,
        end_time=block.end_time,
    ).first()
    if schedule is None:
        schedule = Schedule(
            monitor=monitor,
            weekday=block.weekday,
            start_time=block.start_time,
            end_time=block.end_time,
            is_active=True,
        )
        schedule.full_clean()
        schedule.save()
        result.created += 1
        return
    if not schedule.is_active:
        schedule.is_active = True
        schedule.full_clean()
        schedule.save(update_fields=["is_active", "updated_at"])
        result.reactivated += 1


def _validate_schedule_import_scope(*, actor, monitor: Monitor) -> None:
    if actor is None or actor.role == UserRoleChoices.ADMIN:
        return
    if actor.role != UserRoleChoices.LEADER or not actor.department:
        raise ValidationError("Solo administradores y lideres pueden importar horarios.")
    if monitor.department != actor.department:
        raise ValidationError("Solo puedes importar horarios de tu propia dependencia.")


def _flush_monitor_blocks(
    *,
    monitor_code: Optional[str],
    monitor_name: Optional[str],
    blocks: list[ParsedScheduleBlock],
    result: ScheduleImportResult,
    actor=None,
) -> None:
    if not monitor_code or not blocks:
        return
    monitor = Monitor.objects.filter(codigo_estudiante=str(monitor_code).strip(), is_active=True).first()
    if monitor is None:
        result.missing_monitors.append(f"{monitor_name or 'Sin nombre'} ({monitor_code})")
        return
    try:
        _validate_schedule_import_scope(actor=actor, monitor=monitor)
    except ValidationError:
        result.unauthorized_monitors.append(
            f"{monitor.full_name} ({monitor.codigo_estudiante}) - {monitor.get_department_display()}"
        )
        return
    for block in _merge_schedule_blocks(blocks):
        _save_schedule_block(monitor=monitor, block=block, result=result)
    result.processed_monitors += 1


def _project_lookup() -> dict[str, str]:
    mapping = {"": ""}
    for value, label in PROJECT_CHOICES:
        mapping[value] = value
        mapping[normalize_text(value)] = value
        mapping[normalize_text(label)] = value
    return mapping


def _normalize_project(value: Any) -> str:
    raw_value = str(value or "").strip()
    if not raw_value:
        return ""
    project = _project_lookup().get(raw_value) or _project_lookup().get(normalize_text(raw_value))
    if project is None:
        raise ValidationError("Proyecto curricular no reconocido.")
    return project


def upsert_schedule(
    *,
    monitor,
    weekday: int,
    start_time,
    end_time,
    location: str = "",
    is_active: bool = True,
    asignatura: str = "",
    grupo: str = "",
    docente: str = "",
    proyecto_curricular: str = "",
) -> Schedule:
    schedule = Schedule.objects.filter(
        monitor=monitor,
        weekday=weekday,
        start_time=start_time,
        end_time=end_time,
    ).first()
    return save_schedule(
        instance=schedule,
        monitor=monitor,
        weekday=weekday,
        start_time=start_time,
        end_time=end_time,
        location=location,
        asignatura=asignatura,
        grupo=grupo,
        docente=docente,
        proyecto_curricular=proyecto_curricular,
        is_active=is_active,
    )


def save_schedule(
    *,
    instance: Optional[Schedule] = None,
    monitor: Monitor,
    weekday: int,
    start_time,
    end_time,
    location: str,
    asignatura: str = "",
    grupo: str = "",
    docente: str = "",
    proyecto_curricular: str = "",
    is_active: bool = True,
    actor=None,
) -> Schedule:
    if actor and actor.role != UserRoleChoices.ADMIN and monitor.department != actor.department:
        raise ValidationError("Solo puedes gestionar horarios de tu propia dependencia.")
    schedule = instance or Schedule()
    schedule.monitor = monitor
    schedule.weekday = weekday
    schedule.start_time = start_time
    schedule.end_time = end_time
    schedule.asignatura = str(asignatura or "").strip()
    schedule.grupo = str(grupo or "").strip()
    schedule.docente = str(docente or "").strip()
    schedule.proyecto_curricular = _normalize_project(proyecto_curricular)
    schedule.location = location
    schedule.is_active = is_active
    schedule.full_clean()
    schedule.save()
    return schedule


def delete_schedule(*, schedule: Schedule) -> None:
    schedule.delete()


def _normalize_header_key(value: Any) -> str:
    text = str(value or "").strip().replace("_", " ").replace("-", " ").replace("/", " ")
    return normalize_text(text).rstrip(":")


def _schedule_column_alias_lookup() -> dict[str, str]:
    lookup: dict[str, str] = {}
    for internal_name, aliases in SCHEDULE_ROW_COLUMN_ALIASES.items():
        lookup[_normalize_header_key(internal_name)] = internal_name
        for alias in aliases:
            lookup[_normalize_header_key(alias)] = internal_name
    return lookup


def _schedule_header_map(headers) -> dict[str, int]:
    alias_lookup = _schedule_column_alias_lookup()
    mapped_headers: dict[str, int] = {}
    for index, header in enumerate(headers):
        internal_name = alias_lookup.get(_normalize_header_key(header))
        if internal_name and internal_name not in mapped_headers:
            mapped_headers[internal_name] = index
    return mapped_headers


def _parse_time_value(value: Any) -> time:
    if isinstance(value, time):
        return value
    if isinstance(value, datetime):
        return value.time()
    text = str(value or "").strip()
    for fmt in ("%H:%M:%S", "%H:%M", "%I:%M %p"):
        try:
            return datetime.strptime(text, fmt).time()
        except ValueError:
            continue
    raise ValidationError("Hora invalida. Usa formato HH:MM.")


def _parse_weekday_value(value: Any) -> int:
    if isinstance(value, int) and value in [choice[0] for choice in Schedule.Weekday.choices[:6]]:
        return value
    normalized = normalize_text(str(value or "").strip())
    aliases = {
        "lunes": Schedule.Weekday.MONDAY,
        "monday": Schedule.Weekday.MONDAY,
        "martes": Schedule.Weekday.TUESDAY,
        "tuesday": Schedule.Weekday.TUESDAY,
        "miercoles": Schedule.Weekday.WEDNESDAY,
        "miércoles": Schedule.Weekday.WEDNESDAY,
        "wednesday": Schedule.Weekday.WEDNESDAY,
        "jueves": Schedule.Weekday.THURSDAY,
        "thursday": Schedule.Weekday.THURSDAY,
        "viernes": Schedule.Weekday.FRIDAY,
        "friday": Schedule.Weekday.FRIDAY,
        "sabado": Schedule.Weekday.SATURDAY,
        "sábado": Schedule.Weekday.SATURDAY,
        "saturday": Schedule.Weekday.SATURDAY,
    }
    weekday = aliases.get(normalized)
    if weekday is None:
        raise ValidationError("Dia no reconocido. Usa lunes a sabado.")
    return weekday


def import_schedule_rows_from_workbook(*, uploaded_file, actor=None) -> ScheduleBulkImportResult:
    validate_excel_extension(uploaded_file.name)
    workbook = load_workbook(uploaded_file, read_only=True, data_only=True)
    worksheet = workbook.active
    rows = worksheet.iter_rows(values_only=True)
    try:
        headers = next(rows)
    except StopIteration:
        raise ValidationError("El archivo esta vacio.")

    mapped_headers = _schedule_header_map(headers)
    missing = [column for column in SCHEDULE_ROW_REQUIRED_COLUMNS if column not in mapped_headers]
    if missing:
        raise ValidationError("Faltan columnas requeridas: " + ", ".join(missing))

    result = ScheduleBulkImportResult()
    for row_number, row in enumerate(rows, start=2):
        values = list(row)
        if not any(values):
            continue
        result.total_rows += 1
        monitor_email = str(values[mapped_headers["email_monitor"]] or "").strip().lower()
        try:
            if not monitor_email:
                raise ValidationError("El correo del monitor es obligatorio.")
            monitors = Monitor.objects.select_related("user")
            if actor:
                monitors = visible_monitors_for_user(actor).select_related("user")
            monitor = monitors.filter(user__email__iexact=monitor_email).first()
            if monitor is None:
                result.skipped.append(
                    ScheduleRowIssue(row_number=row_number, monitor_email=monitor_email, reason="Monitor no encontrado.")
                )
                continue
            with transaction.atomic():
                save_schedule(
                    monitor=monitor,
                    weekday=_parse_weekday_value(values[mapped_headers["day"]]),
                    start_time=_parse_time_value(values[mapped_headers["start_time"]]),
                    end_time=_parse_time_value(values[mapped_headers["end_time"]]),
                    location=str(values[mapped_headers["location"]] or "").strip(),
                    asignatura=str(values[mapped_headers["asignatura"]] or "").strip() if "asignatura" in mapped_headers else "",
                    grupo=str(values[mapped_headers["grupo"]] or "").strip() if "grupo" in mapped_headers else "",
                    docente=str(values[mapped_headers["docente"]] or "").strip() if "docente" in mapped_headers else "",
                    proyecto_curricular=values[mapped_headers["proyecto_curricular"]] if "proyecto_curricular" in mapped_headers else "",
                    is_active=True,
                    actor=actor,
                )
            result.created += 1
        except Exception as exc:
            reason = "; ".join(exc.messages) if hasattr(exc, "messages") else str(exc)
            result.errors.append(ScheduleRowIssue(row_number=row_number, monitor_email=monitor_email or "-", reason=reason))
    return result


def _validate_exception_scope(*, actor, department) -> None:
    if actor.role == UserRoleChoices.ADMIN:
        return
    if not actor.department or department != actor.department:
        raise ValidationError("Solo puedes gestionar excepciones de tu propia dependencia.")


def save_schedule_exception(
    *,
    actor,
    instance: Optional[ScheduleException] = None,
    name: str,
    description: str,
    start_date,
    end_date,
    department,
    ignore_lateness: bool,
    approve_overtime: bool,
    is_active: bool,
):
    _validate_exception_scope(actor=actor, department=department)
    exception = instance or ScheduleException()
    previous_state = None
    if instance is not None and instance.pk:
        previous = ScheduleException.objects.get(pk=instance.pk)
        previous_state = {
            "start_date": previous.start_date,
            "end_date": previous.end_date,
            "department": previous.department,
        }

    exception.name = name
    exception.description = description
    exception.start_date = start_date
    exception.end_date = end_date
    exception.department = department
    exception.ignore_lateness = ignore_lateness
    exception.approve_overtime = approve_overtime
    exception.is_active = is_active
    exception.full_clean()
    exception.save()

    updated_sessions = sync_sessions_for_exception_change(
        current_exception=exception,
        previous_start_date=previous_state["start_date"] if previous_state else None,
        previous_end_date=previous_state["end_date"] if previous_state else None,
        previous_department=previous_state["department"] if previous_state else None,
    )
    return exception, updated_sessions


def delete_schedule_exception(*, actor, exception: ScheduleException) -> int:
    _validate_exception_scope(actor=actor, department=exception.department)
    previous_state = {
        "start_date": exception.start_date,
        "end_date": exception.end_date,
        "department": exception.department,
    }
    exception.delete()
    return sync_sessions_for_exception_change(
        previous_start_date=previous_state["start_date"],
        previous_end_date=previous_state["end_date"],
        previous_department=previous_state["department"],
    )


def import_schedules_from_workbook(*, uploaded_file, actor=None) -> ScheduleImportResult:
    validate_excel_extension(uploaded_file.name)
    workbook = load_workbook(uploaded_file, read_only=True, data_only=True)
    worksheet = workbook.active

    result = ScheduleImportResult()
    current_monitor_name: Optional[str] = None
    current_monitor_code: Optional[str] = None
    current_blocks: list[ParsedScheduleBlock] = []
    day_hour_column_index: Optional[int] = None

    for row in worksheet.iter_rows(values_only=True):
        row_values = [str(value).strip() if value is not None else "" for value in row]
        if not any(row_values):
            continue

        monitor_name = _row_label_value_any(row_values, BLOCK_MONITOR_NAME_LABELS)
        if monitor_name is not None:
            _flush_monitor_blocks(
                monitor_code=current_monitor_code,
                monitor_name=current_monitor_name,
                blocks=current_blocks,
                result=result,
                actor=actor,
            )
            current_monitor_name = monitor_name
            current_monitor_code = None
            current_blocks = []
            day_hour_column_index = None
            continue

        monitor_code = _row_label_value_any(row_values, BLOCK_MONITOR_CODE_LABELS)
        if monitor_code is not None:
            current_monitor_code = monitor_code
            continue

        normalized_headers = [_normalize_header_key(value) for value in row_values]
        header_day_hour_column_index = next(
            (index for index, header in enumerate(normalized_headers) if header in BLOCK_DAY_HOUR_HEADER_ALIASES),
            None,
        )
        if header_day_hour_column_index is not None:
            day_hour_column_index = header_day_hour_column_index
            continue

        if day_hour_column_index is None or day_hour_column_index >= len(row_values):
            continue

        day_hour_value = row_values[day_hour_column_index]
        if not day_hour_value:
            continue

        block = _parse_day_hour(day_hour_value)
        if block is None:
            result.skipped_rows += 1
            continue
        current_blocks.append(block)

    _flush_monitor_blocks(
        monitor_code=current_monitor_code,
        monitor_name=current_monitor_name,
        blocks=current_blocks,
        result=result,
        actor=actor,
    )
    return result
