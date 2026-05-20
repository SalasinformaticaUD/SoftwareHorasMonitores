"""Servicios de importacion y conciliacion de asistencia CrossChex.

Este modulo valida archivos Excel, crea trabajos de importacion, transforma
filas en registros crudos, concilia monitores por nombre/dependencia, procesa
sesiones derivadas y permite resolver registros manualmente sin alterar la
marcacion original.
"""

from __future__ import annotations

import logging
from collections import defaultdict
from datetime import datetime, time, timedelta
from typing import Optional, Tuple

from django.core.exceptions import ValidationError
from django.core.files.base import File
from django.db import transaction
from django.db.models import Q
from django.utils import timezone
from openpyxl import load_workbook

from apps.attendance.events import ATTENDANCE_IMPORTED, ATTENDANCE_RECONCILIATION_FAILED
from apps.attendance.models import (
    AttendanceImportJob,
    AttendanceInconsistency,
    AttendanceInconsistencyEvent,
    AttendanceRawRecord,
)
from apps.attendance.validators import (
    coerce_date,
    coerce_datetime,
    coerce_time,
    resolve_attendance_headers,
    resolve_headers,
    validate_excel_extension,
)
from apps.common.choices import (
    AttendanceInconsistencyStatusChoices,
    AttendanceInconsistencyActionChoices,
    AttendanceInconsistencyTypeChoices,
    AttendancePairingStatusChoices,
    DepartmentChoices,
    ImportJobStatusChoices,
    ReconciliationStatusChoices,
    UserRoleChoices,
)
from apps.common.events import DomainEvent, event_bus
from apps.common.permissions import department_allowed
from apps.common.utils import normalize_text
from apps.monitors.models import Monitor

logger = logging.getLogger(__name__)

PAIRING_DUPLICATE_WINDOW = timedelta(minutes=5)
MINIMUM_PAIR_DURATION = timedelta(minutes=30)
JOURNEY_START_TIME = time(hour=6)
JOURNEY_END_TIME = time(hour=22)
JOURNEY_WINDOW_START_TIME = time(hour=5, minute=45)
JOURNEY_WINDOW_END_TIME = time(hour=22, minute=15)

DEPARTMENT_LABELS = {
    DepartmentChoices.PHYSICS: "Monitores Fisica",
    DepartmentChoices.INFORMATICS_LABS: "Monitores Aulas de Software",
    DepartmentChoices.ELECTRICAL: "Monitores Laboratorios",
}

DEPARTMENT_MAPPING = {
    "monitores": "informatics_labs",
    "monitores aulas de software": "informatics_labs",
    "aulas de software": "informatics_labs",
    "monitorias": "informatics_labs",
    "fisica": "physics",
    "monitores fisica": "physics",
    "physics": "physics",
    "monitores aulas de sistemas": "informatics_labs",
    "salas informatica": "informatics_labs",
    "informatica": "informatics_labs",
    "informatics labs": "informatics_labs",
    "monitores laboratorios": "electrical",
    "monitores laboratorio": "electrical",
    "laboratorios": "electrical",
    "electrica": "electrical",
    "electrical": "electrical",
}


def _rewind_uploaded_file(uploaded_file) -> None:
    """Devuelve el puntero de un archivo subido al inicio si es posible.

    Args:
        uploaded_file: Archivo Django o file-like.

    Returns:
        None: Modifica el puntero del archivo en sitio.
    """
    if hasattr(uploaded_file, "seek"):
        uploaded_file.seek(0)


def _department_label(value: str) -> str:
    """Obtiene la etiqueta humana de una dependencia.

    Args:
        value: Valor interno de dependencia.

    Returns:
        str: Etiqueta configurada o texto de respaldo.
    """
    return DEPARTMENT_LABELS.get(value, value or "sin dependencia")


def _match_monitor_queryset(*, raw_full_name: str, raw_department: str):
    """Busca monitores activos que coinciden con nombre y dependencia crudos.

    Args:
        raw_full_name: Nombre recibido desde CrossChex/Excel.
        raw_department: Dependencia recibida desde CrossChex/Excel.

    Returns:
        QuerySet[Monitor]: Candidatos activos compatibles.
    """
    mapped_department = _map_department(normalize_text(raw_department))
    queryset = Monitor.objects.filter(
        normalized_full_name=normalize_text(raw_full_name),
        is_active=True,
    )
    if mapped_department is not None:
        queryset = queryset.filter(department=mapped_department)
    return queryset


def _preview_workbook_departments(uploaded_file) -> tuple[set[str], dict[int, str]]:
    """Lee dependencias de un Excel antes de crear el trabajo de importacion.

    Args:
        uploaded_file: Archivo Excel subido.

    Returns:
        tuple[set[str], dict[int, str]]: Dependencias reconocidas y filas con
        dependencias desconocidas.

    Raises:
        ValidationError: Si el archivo esta vacio o no puede validarse.
    """
    workbook = None
    try:
        _rewind_uploaded_file(uploaded_file)
        workbook = load_workbook(uploaded_file, read_only=True, data_only=True)
        worksheet = workbook.active
        rows = worksheet.iter_rows(values_only=True)
        headers = next(rows, None)
        if not headers:
            raise ValidationError("El archivo esta vacio.")

        header_map = resolve_headers([str(item) for item in headers], required={"department"})
        detected_departments: set[str] = set()
        unknown_rows: dict[int, str] = {}

        for row_number, row in enumerate(rows, start=2):
            if not any(value is not None and str(value).strip() for value in row):
                continue

            raw_department = str(row[header_map["department"]] or "").strip()
            mapped_department = _map_department(normalize_text(raw_department))
            if mapped_department is None:
                unknown_rows[row_number] = raw_department or "(vacio)"
                continue
            detected_departments.add(mapped_department)

        return detected_departments, unknown_rows
    except ValidationError:
        raise
    except Exception as exc:
        raise ValidationError(f"No fue posible validar el archivo Excel antes de subirlo: {exc}")
    finally:
        if workbook is not None:
            workbook.close()
        _rewind_uploaded_file(uploaded_file)

def _validate_import_scope(*, uploaded_file, uploaded_by=None) -> None:
    """Valida que un lider solo importe registros de su dependencia.

    Args:
        uploaded_file: Archivo Excel a validar.
        uploaded_by: Usuario que sube el archivo.

    Returns:
        None: Termina sin error cuando el alcance es valido.

    Raises:
        ValidationError: Si hay dependencias desconocidas o ajenas al lider.
    """
    if uploaded_by is None or uploaded_by.role != UserRoleChoices.LEADER:
        return
    if not uploaded_by.department:
        raise ValidationError("El lider no tiene una dependencia configurada para subir registros.")

    detected_departments, unknown_rows = _preview_workbook_departments(uploaded_file)
    if unknown_rows:
        sample_rows = ", ".join(
            f"fila {row_number}: {value}" for row_number, value in sorted(unknown_rows.items())[:3]
        )
    
        raise ValidationError(
            "No se pudo validar la dependencia del archivo. "
            f"Revisa la columna Departamento. V0alores no reconocidos: {sample_rows}."
        )

    foreign_departments = sorted(
        department for department in detected_departments if department != uploaded_by.department
    )
    if foreign_departments:
        foreign_labels = ", ".join(_department_label(department) for department in foreign_departments)
        raise ValidationError(
            f"Solo puedes subir registros de {_department_label(uploaded_by.department)}. "
            f"El archivo contiene filas de: {foreign_labels}."
        )



def create_import_job(*, uploaded_file, uploaded_by=None) -> AttendanceImportJob:
    """Crea un trabajo de importacion de asistencia.

    Args:
        uploaded_file: Archivo Excel recibido.
        uploaded_by: Usuario que realiza la carga.

    Returns:
        AttendanceImportJob: Trabajo persistido en estado pendiente.

    Raises:
        ValidationError: Si el archivo no es Excel o no respeta el alcance.
    """
    validate_excel_extension(uploaded_file.name)
    _validate_import_scope(uploaded_file=uploaded_file, uploaded_by=uploaded_by)
    return AttendanceImportJob.objects.create(
        uploaded_by=uploaded_by,
        source_file=uploaded_file,
        file_name=uploaded_file.name,
    )


def create_import_job_from_path(*, file_path: str, uploaded_by=None) -> AttendanceImportJob:
    """Crea un trabajo de importacion a partir de una ruta local.

    Args:
        file_path: Ruta del archivo Excel en disco.
        uploaded_by: Usuario asociado a la carga.

    Returns:
        AttendanceImportJob: Trabajo creado con el archivo leido.
    """
    with open(file_path, "rb") as handle:
        django_file = File(handle, name=file_path.split("/")[-1].split("\\")[-1])
        return create_import_job(uploaded_file=django_file, uploaded_by=uploaded_by)


def _serialize_value(value):
    """Serializa valores de Excel para guardarlos en JSON.

    Args:
        value: Valor original de celda.

    Returns:
        object: Valor serializable; fechas se convierten a ISO string.
    """
    if isinstance(value, datetime):
        return value.isoformat()
    if hasattr(value, "isoformat"):
        return value.isoformat()
    return value


def _read_cell(row, header_map: dict[str, int], key: str, default: str = "") -> str:
    """Lee una celda como texto sin fallar si la columna viene vacia."""
    index = header_map.get(key)
    if index is None or index >= len(row):
        return default
    value = row[index]
    if value is None:
        return default
    return str(value).strip()


def _coerce_event_datetime(value) -> datetime:
    """Normaliza una fecha/hora cruda a la zona horaria del proyecto."""
    event_at = coerce_datetime(value)
    if timezone.is_naive(event_at):
        return timezone.make_aware(event_at, timezone.get_current_timezone())
    return event_at


def _existing_raw_record_for_import(
    *,
    raw_full_name: str,
    raw_department: str,
    work_day,
    entry_at,
    exit_at,
) -> Optional[AttendanceRawRecord]:
    """Busca si una fila ya fue importada previamente.

    Args:
        raw_full_name: Nombre crudo.
        raw_department: Dependencia cruda.
        work_day: Fecha del registro.
        entry_at: Hora de entrada.
        exit_at: Hora de salida.

    Returns:
        AttendanceRawRecord | None: Registro existente compatible.
    """
    normalized_department = normalize_text(raw_department)
    department_aliases = _department_aliases(normalized_department)
    return (
        AttendanceRawRecord.objects.filter(
            normalized_full_name=normalize_text(raw_full_name),
            normalized_department__in=department_aliases,
            work_day=work_day,
            entry_at=entry_at,
            exit_at=exit_at,
        )
        .order_by("created_at")
        .first()
    )


def _existing_raw_event_for_import(
    *,
    raw_full_name: str,
    raw_department: str,
    event_at,
    record_type: str,
    operation: str,
    device_number: str,
) -> Optional[AttendanceRawRecord]:
    """Busca una marcacion cruda importada previamente."""
    normalized_department = normalize_text(raw_department)
    department_aliases = _department_aliases(normalized_department)
    return (
        AttendanceRawRecord.objects.filter(
            normalized_full_name=normalize_text(raw_full_name),
            normalized_department__in=department_aliases,
            event_at=event_at,
            record_type=record_type,
            operation=operation,
            device_number=device_number,
        )
        .order_by("created_at")
        .first()
    )


def _match_monitor(raw_record: AttendanceRawRecord) -> Tuple[Optional[Monitor], str]:
    """Resuelve el monitor asociado a un registro crudo.

    Args:
        raw_record: Registro crudo a conciliar.

    Returns:
        tuple[Monitor | None, str]: Monitor encontrado y motivo de fallo cuando
        no hay coincidencia unica.
    """
    mapped_department = _map_department(raw_record.normalized_department)
    if mapped_department is None:
        return None, "La dependencia del registro no coincide con ninguna dependencia configurada."
    matches = _match_monitor_queryset(
        raw_full_name=raw_record.raw_full_name,
        raw_department=raw_record.raw_department,
    )
    if matches.count() == 1:
        return matches.first(), ""
    if matches.count() > 1:
        return None, "Coincidencia ambigua por nombre y dependencia."
    return None, "No se encontró monitor por nombre y dependencia."


def _map_department(normalized_department: str) -> Optional[str]:
    """Mapea una dependencia normalizada al valor interno del sistema.

    Args:
        normalized_department: Texto de dependencia ya normalizado.

    Returns:
        str | None: Valor de ``DepartmentChoices`` o ``None``.
    """
    return DEPARTMENT_MAPPING.get(normalized_department)


def _department_aliases(normalized_department: str) -> set[str]:
    """Obtiene alias equivalentes para detectar duplicados por dependencia.

    Args:
        normalized_department: Dependencia normalizada de la fila.

    Returns:
        set[str]: Alias normalizados que apuntan a la misma dependencia.
    """
    mapped_department = _map_department(normalized_department)
    if mapped_department is None:
        return {normalized_department}
    return {
        alias
        for alias, department in DEPARTMENT_MAPPING.items()
        if department == mapped_department
    }


@transaction.atomic
def reconcile_raw_record(raw_record: AttendanceRawRecord) -> AttendanceRawRecord:
    """Concilia automaticamente un registro crudo con un monitor.

    Args:
        raw_record: Registro pendiente de conciliacion.

    Returns:
        AttendanceRawRecord: Registro actualizado como conciliado o en revision.
    """
    monitor, reason = _match_monitor(raw_record)
    if monitor:
        raw_record.monitor = monitor
        raw_record.reconciliation_status = ReconciliationStatusChoices.MATCHED
        raw_record.manual_review_reason = ""
    else:
        raw_record.reconciliation_status = ReconciliationStatusChoices.MANUAL_REVIEW
        raw_record.manual_review_reason = reason
        event_bus.publish(
            DomainEvent(
                name=ATTENDANCE_RECONCILIATION_FAILED,
                aggregate_id=str(raw_record.id),
                payload={
                    "raw_record_id": str(raw_record.id),
                    "reason": reason,
                    "raw_department": raw_record.raw_department,
                },
            )
        )
    raw_record.save(update_fields=["monitor", "reconciliation_status", "manual_review_reason", "updated_at"])
    return raw_record


@transaction.atomic
def assign_monitor_manually(*, raw_record: AttendanceRawRecord, monitor: Monitor, actor=None) -> AttendanceRawRecord:
    """Asigna manualmente un monitor a un registro crudo.

    Args:
        raw_record: Registro que requiere conciliacion.
        monitor: Monitor seleccionado por el usuario.
        actor: Usuario que ejecuta la asignacion.

    Returns:
        AttendanceRawRecord: Registro marcado como conciliado.

    Raises:
        ValidationError: Si el actor no tiene alcance sobre la dependencia.
    """
    if actor is not None and not department_allowed(actor, monitor.department):
        raise ValidationError("No puedes conciliar monitores de otra dependencia.")
    raw_record.monitor = monitor
    raw_record.reconciliation_status = ReconciliationStatusChoices.MATCHED
    raw_record.manual_review_reason = ""
    raw_record.save(update_fields=["monitor", "reconciliation_status", "manual_review_reason", "updated_at"])
    if raw_record.event_at is not None:
        pair_raw_attendance_events(groups={(monitor.id, raw_record.work_day)})
        process_paired_raw_attendance_events(groups={(monitor.id, raw_record.work_day)})
    elif raw_record.is_processable:
        from apps.work_sessions.services import process_raw_record_to_session

        process_raw_record_to_session(raw_record=raw_record)
    return raw_record


def _event_time(value):
    local_value = timezone.localtime(value) if timezone.is_aware(value) else value
    return local_value.time()


def _is_inside_journey_window(value) -> bool:
    event_time = _event_time(value)
    return JOURNEY_WINDOW_START_TIME <= event_time <= JOURNEY_WINDOW_END_TIME


def _detect_inconsistency(*, raw_record, inconsistency_type: str, message: str) -> AttendanceInconsistency:
    inconsistency, created = AttendanceInconsistency.objects.update_or_create(
        raw_record=raw_record,
        inconsistency_type=inconsistency_type,
        defaults={
            "monitor": raw_record.monitor,
            "work_day": raw_record.work_day,
            "status": AttendanceInconsistencyStatusChoices.PENDING,
            "detected_at": timezone.now(),
            "message": message,
            "validated_by": None,
            "validated_at": None,
            "resolution_note": "",
            "solution_annotation": None,
        },
    )
    AttendanceInconsistencyEvent.objects.create(
        inconsistency=inconsistency,
        action=AttendanceInconsistencyActionChoices.DETECTED,
        note=message if created else f"Detectada nuevamente: {message}",
    )
    return inconsistency


def _add_inconsistency_event(*, inconsistency, action: str, note: str = "", actor=None) -> None:
    AttendanceInconsistencyEvent.objects.create(
        inconsistency=inconsistency,
        actor=actor,
        action=action,
        note=note,
    )


def _auto_resolve_duplicate_inconsistency(*, raw_record, inconsistency: AttendanceInconsistency) -> None:
    raw_record.reconciliation_status = ReconciliationStatusChoices.REJECTED
    raw_record.manual_review_reason = "Invalidado automaticamente por marcacion duplicada dentro de 5 minutos."
    raw_record.processed_at = timezone.now()
    raw_record.processing_error = ""
    raw_record.save(
        update_fields=[
            "reconciliation_status",
            "manual_review_reason",
            "processed_at",
            "processing_error",
            "updated_at",
        ]
    )
    inconsistency.status = AttendanceInconsistencyStatusChoices.RESOLVED
    inconsistency.detected_at = inconsistency.detected_at or timezone.now()
    inconsistency.resolution_note = "Invalidada automaticamente por duplicidad."
    inconsistency.save(update_fields=["status", "detected_at", "resolution_note", "updated_at"])
    _add_inconsistency_event(
        inconsistency=inconsistency,
        action=AttendanceInconsistencyActionChoices.AUTO_RESOLVED,
        note="Marcacion duplicada invalidada automaticamente.",
    )


def _dismiss_open_inconsistencies(*, raw_record) -> None:
    inconsistencies = list(AttendanceInconsistency.objects.filter(
        raw_record=raw_record,
        status__in=[
            AttendanceInconsistencyStatusChoices.PENDING,
            AttendanceInconsistencyStatusChoices.VALIDATED,
        ],
    ))
    AttendanceInconsistency.objects.filter(id__in=[item.id for item in inconsistencies]).update(
        status=AttendanceInconsistencyStatusChoices.DISMISSED,
        resolution_note="Descartada automaticamente por nuevo emparejamiento.",
        updated_at=timezone.now(),
    )
    for inconsistency in inconsistencies:
        _add_inconsistency_event(
            inconsistency=inconsistency,
            action=AttendanceInconsistencyActionChoices.DISMISSED,
            note="Descartada automaticamente por nuevo emparejamiento.",
        )


@transaction.atomic
def validate_attendance_inconsistency(*, inconsistency: AttendanceInconsistency, actor, note: str = ""):
    if inconsistency.monitor_id and not department_allowed(actor, inconsistency.monitor.department):
        raise ValidationError("No puedes validar inconsistencias de otra dependencia.")
    inconsistency.status = AttendanceInconsistencyStatusChoices.VALIDATED
    inconsistency.validated_by = actor
    inconsistency.validated_at = timezone.now()
    inconsistency.resolution_note = (note or "").strip()
    inconsistency.save(update_fields=["status", "validated_by", "validated_at", "resolution_note", "updated_at"])
    _add_inconsistency_event(
        inconsistency=inconsistency,
        actor=actor,
        action=AttendanceInconsistencyActionChoices.ANNOTATION_LINKED,
        note=inconsistency.resolution_note or "Validada manualmente.",
    )
    return inconsistency


@transaction.atomic
def annotate_attendance_inconsistency(
    *,
    inconsistency: AttendanceInconsistency,
    actor,
    annotation_type: str,
    action: str,
    delta_minutes: int,
    description: str,
):
    if inconsistency.monitor_id is None:
        raise ValidationError("La inconsistencia no tiene monitor asociado para crear una anotacion.")
    from apps.annotations.services import create_annotation

    annotation = create_annotation(
        leader=actor,
        monitor=inconsistency.monitor,
        annotation_type=annotation_type,
        description=description,
        action=action,
        delta_minutes=delta_minutes,
        occurred_on=inconsistency.work_day,
    )
    inconsistency.status = AttendanceInconsistencyStatusChoices.RESOLVED
    inconsistency.validated_by = actor
    inconsistency.validated_at = timezone.now()
    inconsistency.resolution_note = description
    inconsistency.save(update_fields=["status", "validated_by", "validated_at", "resolution_note", "updated_at"])
    return annotation


@transaction.atomic
def link_annotation_to_inconsistency(*, inconsistency: AttendanceInconsistency, annotation, actor=None):
    if inconsistency.monitor_id and annotation.monitor_id != inconsistency.monitor_id:
        raise ValidationError("La anotacion no corresponde al monitor de la inconsistencia.")
    inconsistency.solution_annotation = annotation
    inconsistency.status = AttendanceInconsistencyStatusChoices.RESOLVED
    inconsistency.validated_by = actor
    inconsistency.validated_at = timezone.now()
    inconsistency.resolution_note = "Solucion propuesta mediante anotacion administrativa."
    inconsistency.save(
        update_fields=[
            "solution_annotation",
            "status",
            "validated_by",
            "validated_at",
            "resolution_note",
            "updated_at",
        ]
    )
    _add_inconsistency_event(
        inconsistency=inconsistency,
        actor=actor,
        action=AttendanceInconsistencyActionChoices.ANNOTATION_LINKED,
        note=f"Anotacion vinculada: {annotation.id}",
    )
    return inconsistency


@transaction.atomic
def invalidate_inconsistent_raw_record(*, inconsistency: AttendanceInconsistency, actor, reason: str):
    if inconsistency.solution_annotation_id is None:
        raise ValidationError("Antes de invalidar el registro debes registrar una anotacion como solucion.")
    raw_record = reject_raw_record(raw_record=inconsistency.raw_record, actor=actor, reason=reason)
    inconsistency.status = AttendanceInconsistencyStatusChoices.RESOLVED
    inconsistency.validated_by = actor
    inconsistency.validated_at = timezone.now()
    inconsistency.resolution_note = reason.strip()
    inconsistency.save(update_fields=["status", "validated_by", "validated_at", "resolution_note", "updated_at"])
    _add_inconsistency_event(
        inconsistency=inconsistency,
        actor=actor,
        action=AttendanceInconsistencyActionChoices.INVALIDATED,
        note=reason.strip(),
    )
    return raw_record


def _build_pairing_queryset(*, work_day=None, monitor=None, groups=None):
    queryset = AttendanceRawRecord.objects.filter(
        event_at__isnull=False,
        monitor__isnull=False,
        reconciliation_status=ReconciliationStatusChoices.MATCHED,
        processed_at__isnull=True,
    )
    if work_day is not None:
        queryset = queryset.filter(work_day=work_day)
    if monitor is not None:
        queryset = queryset.filter(monitor=monitor)
    if groups:
        group_query = Q()
        for monitor_id, group_work_day in groups:
            if monitor_id is not None and group_work_day is not None:
                group_query |= Q(monitor_id=monitor_id, work_day=group_work_day)
        if not group_query:
            return queryset.none()
        queryset = queryset.filter(group_query)
    return queryset.order_by("monitor_id", "work_day", "event_at", "created_at")


@transaction.atomic
def pair_raw_attendance_events(*, work_day=None, monitor=None, groups=None) -> dict[str, int]:
    """Empareja marcaciones crudas por monitor y dia, ignorando duplicados cercanos."""
    records = list(_build_pairing_queryset(work_day=work_day, monitor=monitor, groups=groups))
    if not records:
        return {"paired": 0, "duplicate_ignored": 0, "unpaired": 0}

    record_ids = [record.id for record in records]
    AttendanceRawRecord.objects.filter(id__in=record_ids).update(
        pairing_status=AttendancePairingStatusChoices.PENDING,
        paired_record=None,
        duplicate_of=None,
        paired_at=None,
        pairing_reason="",
        entry_at=None,
        exit_at=None,
    )
    AttendanceInconsistency.objects.filter(
        raw_record_id__in=record_ids,
        status__in=[
            AttendanceInconsistencyStatusChoices.PENDING,
            AttendanceInconsistencyStatusChoices.VALIDATED,
        ],
    ).update(
        status=AttendanceInconsistencyStatusChoices.DISMISSED,
        resolution_note="Descartada automaticamente por nuevo emparejamiento.",
        updated_at=timezone.now(),
    )

    grouped_records: dict[tuple[str, object], list[AttendanceRawRecord]] = defaultdict(list)
    for record in records:
        grouped_records[(str(record.monitor_id), record.work_day)].append(record)

    counts = {"paired": 0, "duplicate_ignored": 0, "unpaired": 0}
    paired_at = timezone.now()
    for group_records in grouped_records.values():
        valid_records = []
        last_valid = None
        for record in group_records:
            if not _is_inside_journey_window(record.event_at):
                record.pairing_status = AttendancePairingStatusChoices.UNPAIRED
                record.duplicate_of = None
                record.paired_record = None
                record.paired_at = paired_at
                record.pairing_reason = "Marcacion fuera de la ventana permitida de 05:45 a 22:15."
                record.entry_at = None
                record.exit_at = None
                record.save(
                    update_fields=[
                        "pairing_status",
                        "duplicate_of",
                        "paired_record",
                        "paired_at",
                        "pairing_reason",
                        "entry_at",
                        "exit_at",
                        "updated_at",
                    ]
                )
                _detect_inconsistency(
                    raw_record=record,
                    inconsistency_type=AttendanceInconsistencyTypeChoices.OUT_OF_DAY_WINDOW,
                    message="Marcacion fuera de la ventana permitida de 05:45 a 22:15.",
                )
                counts["unpaired"] += 1
                continue

            if last_valid is not None and record.event_at < last_valid.event_at + PAIRING_DUPLICATE_WINDOW:
                record.pairing_status = AttendancePairingStatusChoices.DUPLICATE_IGNORED
                record.duplicate_of = last_valid
                record.paired_record = None
                record.paired_at = paired_at
                record.pairing_reason = "Marcacion ignorada por repetirse dentro de una ventana de 5 minutos."
                record.entry_at = None
                record.exit_at = None
                record.save(
                    update_fields=[
                        "pairing_status",
                        "duplicate_of",
                        "paired_record",
                        "paired_at",
                        "pairing_reason",
                        "entry_at",
                        "exit_at",
                        "updated_at",
                    ]
                )
                duplicate_inconsistency = _detect_inconsistency(
                    raw_record=record,
                    inconsistency_type=AttendanceInconsistencyTypeChoices.DUPLICATE_MARK,
                    message="Marcacion repetida dentro de la ventana de 5 minutos.",
                )
                _auto_resolve_duplicate_inconsistency(raw_record=record, inconsistency=duplicate_inconsistency)
                counts["duplicate_ignored"] += 1
                continue
            valid_records.append(record)
            last_valid = record

        pairable_count = len(valid_records) - (len(valid_records) % 2)
        for index in range(0, pairable_count, 2):
            entry_record = valid_records[index]
            exit_record = valid_records[index + 1]
            entry_record.pairing_status = AttendancePairingStatusChoices.PAIRED
            entry_record.paired_record = exit_record
            entry_record.duplicate_of = None
            entry_record.paired_at = paired_at
            entry_record.pairing_reason = ""
            entry_record.entry_at = _event_time(entry_record.event_at)
            entry_record.exit_at = _event_time(exit_record.event_at)
            entry_record.save(
                update_fields=[
                    "pairing_status",
                    "paired_record",
                    "duplicate_of",
                    "paired_at",
                    "pairing_reason",
                    "entry_at",
                    "exit_at",
                    "updated_at",
                ]
            )

            exit_record.pairing_status = AttendancePairingStatusChoices.PAIRED
            exit_record.paired_record = entry_record
            exit_record.duplicate_of = None
            exit_record.paired_at = paired_at
            exit_record.pairing_reason = ""
            exit_record.entry_at = None
            exit_record.exit_at = None
            exit_record.save(
                update_fields=[
                    "pairing_status",
                    "paired_record",
                    "duplicate_of",
                    "paired_at",
                    "pairing_reason",
                    "entry_at",
                    "exit_at",
                    "updated_at",
                ]
            )
            if exit_record.event_at - entry_record.event_at < MINIMUM_PAIR_DURATION:
                entry_record.pairing_reason = "Emparejamiento menor a 30 minutos pendiente de revision."
                entry_record.save(update_fields=["pairing_reason", "updated_at"])
                _detect_inconsistency(
                    raw_record=entry_record,
                    inconsistency_type=AttendanceInconsistencyTypeChoices.SHORT_PAIR,
                    message="La pareja de marcaciones tiene una duracion menor a 30 minutos.",
                )
            counts["paired"] += 1

        if len(valid_records) % 2:
            unpaired_record = valid_records[-1]
            unpaired_record.pairing_status = AttendancePairingStatusChoices.UNPAIRED
            unpaired_record.paired_record = None
            unpaired_record.duplicate_of = None
            unpaired_record.paired_at = paired_at
            unpaired_record.pairing_reason = "Marcacion sin pareja del mismo dia."
            unpaired_record.entry_at = None
            unpaired_record.exit_at = None
            unpaired_record.save(
                update_fields=[
                    "pairing_status",
                    "paired_record",
                    "duplicate_of",
                    "paired_at",
                    "pairing_reason",
                    "entry_at",
                    "exit_at",
                    "updated_at",
                ]
            )
            inconsistency_type = AttendanceInconsistencyTypeChoices.END_OF_DAY
            message = "Entrada sin salida al cierre de jornada de las 22:00."
            if _event_time(unpaired_record.event_at) >= JOURNEY_END_TIME:
                inconsistency_type = AttendanceInconsistencyTypeChoices.ODD_MARK
                message = "Marcacion impar sin pareja dentro del mismo dia."
            _detect_inconsistency(
                raw_record=unpaired_record,
                inconsistency_type=inconsistency_type,
                message=message,
            )
            counts["unpaired"] += 1

    return counts


@transaction.atomic
def process_paired_raw_attendance_events(*, work_day=None, monitor=None, groups=None) -> dict[str, int]:
    """Crea sesiones calculadas a partir de pares validos de marcaciones crudas."""
    queryset = AttendanceRawRecord.objects.filter(
        event_at__isnull=False,
        monitor__isnull=False,
        reconciliation_status=ReconciliationStatusChoices.MATCHED,
        pairing_status=AttendancePairingStatusChoices.PAIRED,
        entry_at__isnull=False,
        exit_at__isnull=False,
        processed_at__isnull=True,
    ).exclude(
        inconsistencies__status__in=[
            AttendanceInconsistencyStatusChoices.PENDING,
            AttendanceInconsistencyStatusChoices.VALIDATED,
        ]
    ).select_related("monitor").distinct()
    if work_day is not None:
        queryset = queryset.filter(work_day=work_day)
    if monitor is not None:
        queryset = queryset.filter(monitor=monitor)
    if groups:
        group_query = Q()
        for monitor_id, group_work_day in groups:
            if monitor_id is not None and group_work_day is not None:
                group_query |= Q(monitor_id=monitor_id, work_day=group_work_day)
        if not group_query:
            return {"processed": 0, "failed": 0}
        queryset = queryset.filter(group_query)

    from apps.work_sessions.services import process_raw_record_to_session

    processed = 0
    failed = 0
    for raw_record in queryset.order_by("monitor_id", "work_day", "entry_at"):
        try:
            process_raw_record_to_session(raw_record=raw_record)
            raw_record.refresh_from_db(fields=["processed_at"])
            if raw_record.paired_record_id and raw_record.paired_record.processed_at is None:
                raw_record.paired_record.processed_at = raw_record.processed_at
                raw_record.paired_record.processing_error = ""
                raw_record.paired_record.save(update_fields=["processed_at", "processing_error", "updated_at"])
            processed += 1
        except Exception as exc:  # pragma: no cover - defensive path logged for operational review
            raw_record.processing_error = str(exc)
            raw_record.save(update_fields=["processing_error", "updated_at"])
            failed += 1
            logger.exception("attendance_pair_processing_failed raw_record=%s error=%s", raw_record.id, exc)
    return {"processed": processed, "failed": failed}


def import_workbook(job: AttendanceImportJob) -> AttendanceImportJob:
    """Procesa un trabajo de importacion de asistencia desde Excel.

    Args:
        job: Trabajo con archivo fuente asociado.

    Returns:
        AttendanceImportJob: Trabajo actualizado con totales y estado final.

    Raises:
        Exception: Propaga errores globales que impiden procesar el archivo.
    """
    job.status = ImportJobStatusChoices.PROCESSING
    job.started_at = timezone.now()
    job.error_message = ""
    job.save(update_fields=["status", "started_at", "error_message", "updated_at"])

    try:
        workbook = load_workbook(job.source_file.path, read_only=True, data_only=True)
        worksheet = workbook.active
        rows = worksheet.iter_rows(values_only=True)
        headers = next(rows, None)
        if not headers:
            raise ValueError("El archivo está vacío.")
        attendance_format, header_map = resolve_attendance_headers([str(item) for item in headers])

        imported_rows = 0
        failed_rows = 0
        total_rows = 0
        affected_pairing_groups = set()

        for row_number, row in enumerate(rows, start=2):
            if not any(value is not None and str(value).strip() for value in row):
                continue
            total_rows += 1
            payload = {str(headers[index]): _serialize_value(value) for index, value in enumerate(row)}
            try:
                raw_full_name = _read_cell(row, header_map, "full_name")
                raw_department = _read_cell(row, header_map, "department")
                raw_user_number = _read_cell(row, header_map, "num_user")
                raw_user_id = _read_cell(row, header_map, "id_user")

                if attendance_format == "raw":
                    event_at = _coerce_event_datetime(row[header_map["event_at"]])
                    local_event_at = timezone.localtime(event_at) if timezone.is_aware(event_at) else event_at
                    record_type = _read_cell(row, header_map, "record_type")
                    operation = _read_cell(row, header_map, "operation")
                    device_number = _read_cell(row, header_map, "device_number")
                    existing_raw_record = _existing_raw_event_for_import(
                        raw_full_name=raw_full_name,
                        raw_department=raw_department,
                        event_at=event_at,
                        record_type=record_type,
                        operation=operation,
                        device_number=device_number,
                    )
                    if existing_raw_record is not None:
                        logger.info(
                            "attendance_import_row_skipped_duplicate row=%s existing_raw_record=%s",
                            row_number,
                            existing_raw_record.id,
                        )
                        continue
                    matched_monitors = _match_monitor_queryset(
                        raw_full_name=raw_full_name,
                        raw_department=raw_department,
                    )
                    monitor = matched_monitors.first() if matched_monitors.count() == 1 else None
                    if monitor is not None:
                        affected_pairing_groups.add((monitor.id, local_event_at.date()))
                    raw_record = AttendanceRawRecord.objects.create(
                        import_job=job,
                        row_number=row_number,
                        raw_full_name=raw_full_name,
                        raw_department=raw_department,
                        raw_user_number=raw_user_number,
                        raw_user_id=raw_user_id,
                        work_day=local_event_at.date(),
                        event_at=event_at,
                        record_type=record_type,
                        operation=operation,
                        exception_description=_read_cell(row, header_map, "description"),
                        shift=_read_cell(row, header_map, "shift"),
                        identification_code=_read_cell(row, header_map, "identification_code"),
                        identification=_read_cell(row, header_map, "identification"),
                        task_code=_read_cell(row, header_map, "task_code"),
                        device_number=device_number,
                        marked=_read_cell(row, header_map, "marked"),
                        monitor=monitor,
                        raw_payload=payload,
                    )
                    reconcile_raw_record(raw_record)
                else:
                    worked_time = coerce_time(row[header_map["worked_time"]]) if "worked_time" in header_map else None
                    entry_at = coerce_datetime(row[header_map["entry_at"]])
                    work_day = f"{entry_at.year}-{entry_at.month}-{entry_at.day}"
                    entry_at = f"{entry_at.hour}:{entry_at.minute}:{entry_at.second}"
                    exit_at = coerce_datetime(
                        row[header_map["exit_at"]],
                        fallback_date=work_day or entry_at.date(),
                    )
                    exit_at = f"{exit_at.hour}:{exit_at.minute}:{exit_at.second}"
                    work_day = coerce_date(work_day)
                    entry_at = coerce_time(entry_at)
                    exit_at = coerce_time(exit_at)
                    existing_raw_record = _existing_raw_record_for_import(
                        raw_full_name=raw_full_name,
                        raw_department=raw_department,
                        work_day=work_day,
                        entry_at=entry_at,
                        exit_at=exit_at,
                    )
                    if existing_raw_record is not None:
                        logger.info(
                            "attendance_import_row_skipped_duplicate row=%s existing_raw_record=%s",
                            row_number,
                            existing_raw_record.id,
                        )
                        continue
                    matched_monitors = _match_monitor_queryset(
                        raw_full_name=raw_full_name,
                        raw_department=raw_department,
                    )
                    monitor = matched_monitors.first() if matched_monitors.count() == 1 else None
                    raw_record = AttendanceRawRecord.objects.create(
                        import_job=job,
                        row_number=row_number,
                        raw_full_name=raw_full_name,
                        raw_department=raw_department,
                        raw_user_number=raw_user_number,
                        raw_user_id=raw_user_id,
                        work_day=work_day,
                        entry_at=entry_at,
                        exit_at=exit_at,
                        worked_time=worked_time,
                        exception_description=_read_cell(row, header_map, "description"),
                        monitor=monitor,
                        raw_payload=payload,
                    )
                    reconcile_raw_record(raw_record)
                    if raw_record.is_processable:
                        from apps.work_sessions.services import process_raw_record_to_session
                        process_raw_record_to_session(raw_record=raw_record)
                imported_rows += 1
            except Exception as exc:  # pragma: no cover - guarded by tests around service output
                failed_rows += 1
                logger.exception("attendance_import_row_failed row=%s error=%s", row_number, exc)

        if attendance_format == "raw" and affected_pairing_groups:
            pair_raw_attendance_events(groups=affected_pairing_groups)
            process_paired_raw_attendance_events(groups=affected_pairing_groups)

        job.total_rows = total_rows
        job.imported_rows = imported_rows
        job.failed_rows = failed_rows
        job.status = ImportJobStatusChoices.COMPLETED
        job.finished_at = timezone.now()
        job.save(
            update_fields=[
                "total_rows",
                "imported_rows",
                "failed_rows",
                "status",
                "finished_at",
                "updated_at",
            ]
        )
        event_bus.publish(
            DomainEvent(
                name=ATTENDANCE_IMPORTED,
                aggregate_id=str(job.id),
                payload={
                    "job_id": str(job.id),
                    "file_name": job.file_name,
                    "imported_rows": imported_rows,
                    "failed_rows": failed_rows,
                },
            )
        )
        return job
    except Exception as exc:
        job.status = ImportJobStatusChoices.FAILED
        job.finished_at = timezone.now()
        job.error_message = str(exc)
        job.save(update_fields=["status", "finished_at", "error_message", "updated_at"])
        logger.exception("attendance_import_failed job=%s error=%s", job.id, exc)
        raise

@transaction.atomic
def reject_raw_record(*, raw_record: AttendanceRawRecord, actor, reason: str) -> AttendanceRawRecord:
    """Rechaza administrativamente un registro crudo sin tocar su payload.

    Args:
        raw_record: Registro crudo a rechazar.
        actor: Lider o administrador que realiza el rechazo.
        reason: Motivo obligatorio de rechazo.

    Returns:
        AttendanceRawRecord: Registro marcado como rechazado.

    Raises:
        ValidationError: Si el actor no tiene alcance, falta motivo o el
        registro ya genero una sesion que debe invalidarse aparte.
    """
    mapped_department = _map_department(raw_record.normalized_department)
    if raw_record.monitor_id and not department_allowed(actor, raw_record.monitor.department):
        raise ValidationError("No puedes rechazar registros de otra dependencia.")
    if mapped_department and not department_allowed(actor, mapped_department):
        raise ValidationError("No puedes rechazar registros de otra dependencia.")
    reason = (reason or "").strip()
    if not reason:
        raise ValidationError("Rechazar un registro requiere un motivo.")
    if hasattr(raw_record, "work_session"):
        raise ValidationError("Este registro ya tiene una sesion procesada; invalida la sesion derivada.")
    raw_record.reconciliation_status = ReconciliationStatusChoices.REJECTED
    raw_record.manual_review_reason = f"Rechazado administrativamente: {reason}"
    raw_record.save(update_fields=["reconciliation_status", "manual_review_reason", "updated_at"])
    return raw_record
